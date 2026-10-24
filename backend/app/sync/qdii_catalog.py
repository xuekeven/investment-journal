"""QDII identities from the official CSRC public fund catalogue."""

import io
from dataclasses import dataclass
from datetime import date

import httpx
from openpyxl import load_workbook


@dataclass(frozen=True)
class QdiiCatalog:
    names: dict[str, str]
    snapshot_date: date
    source_url: str


def fetch_qdii_catalog() -> QdiiCatalog:
    # Imported at call time because the CSRC collector also uses SSE master data.
    from app.sync.csrc_funds import fetch_product_index, snapshot_date_from_url

    with httpx.Client(
        timeout=httpx.Timeout(40, connect=10),
        headers={"User-Agent": "investment-journal/0.1"},
        transport=httpx.HTTPTransport(retries=2),
    ) as client:
        content, source_url = fetch_product_index(client)
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    snapshot_date = snapshot_date_from_url(source_url)
    names: dict[str, str] = {}
    for row in workbook.active.iter_rows(min_row=3, values_only=True):
        code = str(row[1] or "").strip().zfill(6)
        name = str(row[2] or "").strip()
        if code.isdigit() and len(code) == 6 and "QDII" in name.upper():
            names[code] = name
    return QdiiCatalog(names, snapshot_date, source_url)


def fetch_qdii_catalog_names() -> dict[str, str]:
    return fetch_qdii_catalog().names
